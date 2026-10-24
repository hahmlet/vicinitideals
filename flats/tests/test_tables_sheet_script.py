"""A table corrected in a spreadsheet comes back as cells (scripts/flats_tables_sheet.py).

What a person does in the sheet -- type a cell's words, merge cells the page
prints as one, bold a header -- must come back as exactly that cell, span and
header; the checks in flats.provenance.tables then judge the words.
"""

from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "flats_tables_sheet.py"
spec = importlib.util.spec_from_file_location("flats_tables_sheet", SCRIPT)
sheet_script = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sheet_script)


def test_merged_cells_become_spans_and_bold_cells_headers():
    sheet = Workbook().active
    top = sheet_script.FIRST_ROW
    sheet["A1"] = sheet_script.HOW  # above the grid: not a cell
    sheet.cell(row=top, column=1, value="USE CATEGORY").font = Font(bold=True)
    sheet.cell(row=top, column=2, value="STATUS").font = Font(bold=True)
    sheet.cell(row=top + 1, column=1, value="RESIDENTIAL USE CATEGORIES")
    sheet.merge_cells(start_row=top + 1, start_column=1, end_row=top + 1, end_column=2)
    sheet.cell(row=top + 2, column=1, value="Household\nLiving")
    sheet.cell(row=top + 2, column=2, value="P (L)")
    sheet.cell(row=top + 3, column=2, value="P")  # column 1 left empty: no cell

    rows, cols, cells = sheet_script._grid(sheet)

    assert (rows, cols) == (4, 2)
    by_text = {c["text"]: c for c in cells}
    assert set(by_text) == {"USE CATEGORY", "STATUS", "RESIDENTIAL USE CATEGORIES", "Household Living", "P (L)", "P"}
    assert by_text["USE CATEGORY"]["header"] and not by_text["P (L)"]["header"]
    assert by_text["RESIDENTIAL USE CATEGORIES"]["row"] == [1, 2]
    assert by_text["RESIDENTIAL USE CATEGORIES"]["col"] == [0, 2]
    assert by_text["P"]["row"] == [3, 4] and by_text["P"]["col"] == [1, 2]
    assert all(c["box"] is None for c in cells)


def test_a_sheet_left_as_exported_is_the_same_table_and_a_changed_one_is_not():
    base = {
        "cells": [
            {"row": [0, 1], "col": [0, 1], "text": "USE\nCATEGORY", "header": True, "box": [0, 0, 1, 1]},
            {"row": [0, 1], "col": [1, 2], "text": "STATUS", "header": True, "box": [1, 0, 2, 1]},
            {"row": [1, 2], "col": [0, 1], "text": "", "header": False, "box": None},
        ]
    }
    same = [
        {"row": [0, 1], "col": [0, 1], "text": "USE CATEGORY", "header": True, "box": None},
        {"row": [0, 1], "col": [1, 2], "text": "STATUS", "header": True, "box": None},
    ]
    assert sheet_script._same(base, same)
    assert not sheet_script._same(base, [same[0], {**same[1], "header": False}])
    assert not sheet_script._same(base, [same[0], {**same[1], "col": [1, 3]}])
