"""Sources & Uses per-project totals row sits below the LONGER list.

Regression for the bug the full-workbook error sweep caught on
2026-09-28: each project block writes Use lines down columns A-B and
its capital sources down columns D-E, side by side, then one shared
row carrying "Total Uses P{n}" and "Total Sources P{n}". That row used
to be placed under the Use lines only. A project with more sources
than Use rows had a source overwritten by the total, whose formula then
summed itself (E4 = E4+E5 ...) -- a circular reference that recalcs to
an error.

The fixture: one Use line (3 rows with its category header and
subtotal) and four capital sources (4 rows).
"""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.exporters.investor_export import export_investor_workbook
from app.models.capital import (
    CapitalModule,
    CapitalModuleProject,
    EquityRole,
    VehicleType,
)
from app.models.deal import UseLine
from app.models.project import Project
from tests.conftest import (
    seed_deal_model_with_financials,
    seed_opportunity,
    seed_org,
)
from tests.exporters._parity_helpers import (
    RecalcUnavailableError,
    recalc_workbook,
)

SHEET = "Sources & Uses"
USE_AMOUNT = Decimal("1000000")
SOURCE_AMOUNTS = [
    Decimal("500000"), Decimal("200000"), Decimal("150000"), Decimal("150000"),
]


async def _seed(session: AsyncSession):
    org, user = await seed_org(session)
    opportunity = await seed_opportunity(
        session, org, user, name="S&U more sources than uses"
    )
    deal_model, _, _, _ = await seed_deal_model_with_financials(
        session, opportunity, user
    )
    project = (
        await session.execute(
            select(Project).where(Project.scenario_id == deal_model.id)
        )
    ).scalar_one()
    session.add(UseLine(
        project_id=project.id, label="Purchase Price", amount=USE_AMOUNT,
        phase="acquisition", cost_category="acquisition",
    ))
    modules = []
    for i, amt in enumerate(SOURCE_AMOUNTS, start=1):
        is_debt = i == 1
        m = CapitalModule(
            scenario_id=deal_model.id,
            label=f"Source {i}",
            vehicle_type=(VehicleType.debt if is_debt else VehicleType.equity).value,
            equity_role=None if is_debt else EquityRole.lp.value,
            stack_position=i,
            source=(
                {"amount": str(amt), "interest_rate_pct": 6.5,
                 "amort_term_years": 30, "hold_term_years": 10}
                if is_debt else {"amount": str(amt)}
            ),
            carry=(
                {"carry_type": "pi", "payment_frequency": "monthly"}
                if is_debt else
                {"carry_type": "none", "payment_frequency": "monthly"}
            ),
            exit_terms={"exit_type": "full_payoff", "trigger": "sale"},
            active_phase_start="acquisition", active_phase_end="exit",
        )
        modules.append(m)
    session.add_all(modules)
    await session.flush()
    session.add_all([
        CapitalModuleProject(
            capital_module_id=m.id, project_id=project.id, amount=amt,
        )
        for m, amt in zip(modules, SOURCE_AMOUNTS)
    ])
    await session.commit()
    return deal_model


def _find(ws, col: int, text: str) -> int:
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.strip() == text:
            return r
    raise AssertionError(f"{text!r} not found in column {col} of {ws.title}")


async def test_project_totals_row_sits_below_every_source(
    session: AsyncSession, tmp_path: Path,
):
    scenario = await _seed(session)
    path = tmp_path / "wb.xlsx"
    path.write_bytes(await export_investor_workbook(scenario.id, session))

    ws = load_workbook(path)[SHEET]
    total_row = _find(ws, 4, "Total Sources P1")
    # Both totals share one row.
    assert ws.cell(row=total_row, column=1).value == "Total Uses P1"

    source_rows = [
        r for r in range(1, total_row)
        if isinstance(ws.cell(row=r, column=4).value, str)
        and ws.cell(row=r, column=4).value.startswith("Source ")
    ]
    assert len(source_rows) == len(SOURCE_AMOUNTS), (
        f"a source row was overwritten: found {source_rows}"
    )
    formula = ws.cell(row=total_row, column=5).value
    assert formula == "=" + "+".join(f"E{r}" for r in source_rows), formula
    assert f"E{total_row}" not in formula.replace("=", "+").split("+"), (
        f"Total Sources P1 references itself: {formula}"
    )


async def test_project_totals_evaluate_to_the_right_numbers(
    session: AsyncSession, tmp_path: Path,
):
    scenario = await _seed(session)
    path = tmp_path / "wb.xlsx"
    path.write_bytes(await export_investor_workbook(scenario.id, session))
    total_row = _find(load_workbook(path)[SHEET], 4, "Total Sources P1")

    try:
        recalc_workbook(path)
    except RecalcUnavailableError as exc:
        pytest.skip(f"no recalc backend available: {exc}")

    ws = load_workbook(path, data_only=True)[SHEET]
    total_sources = ws.cell(row=total_row, column=5).value
    total_uses = ws.cell(row=total_row, column=2).value
    assert isinstance(total_sources, (int, float)), total_sources
    assert isinstance(total_uses, (int, float)), total_uses
    assert total_sources == pytest.approx(float(sum(SOURCE_AMOUNTS)), abs=0.01)
    assert total_uses == pytest.approx(float(USE_AMOUNT), abs=0.01)
